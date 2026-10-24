import re
import json
import asyncio
import hashlib
from datetime import date
from pathlib import Path
from typing import Dict, List, Optional

from pydantic import BaseModel

from src.config import settings
from src.utils.logger import custom_print


def file_purpose(text: str) -> str:
    t = (text or "").lower()
    if re.search(r"cover\s*letter|covering letter|motivation letter|letter of (interest|motivation)", t):
        return "cover_letter"
    if re.search(r"resume|résumé|\bcv\b|curriculum", t):
        return "cv"
    if re.search(r"portfolio|work sample|transcript|certificate|right to work|passport|visa|licen[cs]e|id document|photo|video|reference", t):
        return "other"
    return "unknown"


def safe_name(s: str) -> str:
    return re.sub(r"[^\w\-]+", "_", s or "").strip("_")[:60] or "job"


def cv_path(lane: Optional[str]) -> Optional[Path]:
    p = settings.cv_for_lane(lane)
    if not p:
        custom_print("WARNING", f"No CV file found for lane '{lane}' - check CV_FILES in config.py")
    return p


_CV_TEXT_CACHE = {}


def cv_text(lane: Optional[str]) -> str:
    p = settings.cv_for_lane(lane)
    if not p:
        return ""
    if str(p) in _CV_TEXT_CACHE:
        return _CV_TEXT_CACHE[str(p)]
    text = ""
    try:
        if p.suffix.lower() == ".docx":
            from docx import Document
            d = Document(str(p))
            parts = [para.text for para in d.paragraphs if para.text.strip()]
            for t in d.tables:
                for row in t.rows:
                    parts.append(" | ".join(c.text.strip() for c in row.cells if c.text.strip()))
            text = "\n".join(parts)
        elif p.suffix.lower() == ".pdf":
            from pypdf import PdfReader
            text = "\n".join(pg.extract_text() or "" for pg in PdfReader(str(p)).pages)
        elif p.suffix.lower() in (".txt", ".md"):
            text = p.read_text(encoding="utf-8", errors="ignore")
    except Exception as e:
        custom_print("WARNING", f"Could not read CV text from {p.name}: {e}")
    _CV_TEXT_CACHE[str(p)] = text
    return text


def write_cover_letter(job: dict, text: str) -> Optional[Path]:
    settings.COVER_LETTER_DIR.mkdir(parents=True, exist_ok=True)
    stem = f"Simon_Che_de_Boer_Cover_Letter_{safe_name(job.get('company', ''))}_{safe_name(str(job.get('linkedin_job_id') or job.get('id') or ''))}"
    try:
        from docx import Document
        from docx.shared import Pt
        doc = Document()
        style = doc.styles["Normal"]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        doc.add_paragraph(date.today().strftime("%d %B %Y"))
        for para in [p.strip() for p in text.split("\n") if p.strip()]:
            doc.add_paragraph(para)
        out = settings.COVER_LETTER_DIR / f"{stem}.docx"
        doc.save(str(out))
        return out
    except ImportError:
        out = settings.COVER_LETTER_DIR / f"{stem}.txt"
        out.write_text(text, encoding="utf-8")
        custom_print("WARNING", "python-docx not installed - wrote cover letter as .txt (pip install python-docx)")
        return out


# ── Tailored CV per job type ──
class _Role(BaseModel):
    title: str
    company: str
    dates: str
    intro: str = ""
    bullets: List[str] = []


class _SkillGroup(BaseModel):
    name: str
    items: List[str]


class CVVariant(BaseModel):
    headline: str
    summary: str
    specialisations: List[str]
    experience: List[_Role]
    skills: List[_SkillGroup]
    recognition: List[str]


CV_VARIANT_VERSION = "1"    # bump to rebuild every variant after changing the prompt or layout
CV_FILE_NAME = "Simon_Che_de_Boer_CV.docx"


def _cv_sources() -> str:
    """Everything a variant may say: both uploaded CVs plus the structured candidate data. Nothing else."""
    from src.services.cv_data import CV_SUPPLEMENTARY, APPLICATION_PROFILE
    sup = {k: CV_SUPPLEMENTARY[k] for k in ("current_work", "working_style", "technology_experience_years") if k in CV_SUPPLEMENTARY}
    cvs = []
    for lane in ("applied_ai", "platform_arch"):
        p = settings.CV_FILES.get(lane)
        if p and (settings.BASE_DIR / p).exists():
            cvs.append(_read_cv(settings.BASE_DIR / p))
    edu = APPLICATION_PROFILE.get("education", [])
    blocks = [f"=== CV {'AB'[i]} ===\n{t}" for i, t in enumerate(cvs[:2])]
    blocks += ["=== CANDIDATE DATA ===\n" + json.dumps(sup, indent=1, ensure_ascii=False), "=== EDUCATION ===\n" + json.dumps(edu, ensure_ascii=False)]
    facts = settings.CV_VARIANT_DIR / "career_facts.md"      # full verified career history (private, git-ignored)
    if facts.exists():
        blocks.append("=== CAREER FACTS ===\n" + facts.read_text(encoding="utf-8"))
    return "\n\n".join(blocks)


def _read_cv(p: Path) -> str:
    from docx import Document
    d = Document(str(p))
    return "\n".join(x.text for x in d.paragraphs if x.text.strip())


def _unsupported(cv: Dict, source: str) -> List[str]:
    """Numbers and companies in the variant that do not appear in the sources."""
    src = source.lower().replace(",", "")
    parts = [cv["headline"], cv["summary"], *cv["specialisations"], *cv["recognition"],
             *[f'{r["title"]} {r["intro"]} {" ".join(r["bullets"])}' for r in cv["experience"]],
             *[" ".join(g["items"]) for g in cv["skills"]]]
    issues = []
    for part in parts:
        for num in re.findall(r"\d[\d,.]*", part):
            n = num.replace(",", "").rstrip(".")
            if n and n not in src:
                issues.append(f"'{num}' (in: {part[:70]})")
    for r in cv["experience"]:
        company = re.split(r"[(|,]", r["company"])[0].strip().lower()
        if company and company not in src:
            issues.append(f"company '{r['company']}'")
        for y in re.findall(r"(?:19|20)\d\d", r["dates"]):
            if y not in src:
                issues.append(f"date '{r['dates']}' for {r['company']}")
    return issues


def _render_cv(cv: Dict, header: List[str], education: List[Dict], links: List[str], out: Path):
    from docx import Document
    from docx.shared import Pt, Inches
    from src.services.form_engine.answerer import humanize
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(0.7)
        s.top_margin = s.bottom_margin = Inches(0.6)
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(10.5)
    st.paragraph_format.space_after = Pt(3)

    def para(text, bold=False, size=None, italic=False, style=None):
        p = doc.add_paragraph(style=style)
        r = p.add_run(text)
        r.bold, r.italic = bold, italic
        if size:
            r.font.size = Pt(size)
        return p

    def heading(text):
        p = para(text.upper(), bold=True, size=11)
        p.paragraph_format.space_before = Pt(10)

    para(header[0], bold=True, size=18)
    para(humanize(cv["headline"]), bold=True, size=11)
    para(header[1])
    heading("Professional Summary")
    para(humanize(cv["summary"]))
    heading("Core Specialisations")
    for s in cv["specialisations"]:
        para(humanize(s), style="List Bullet")
    heading("Experience")
    for r in cv["experience"]:
        p = para(f'{r["title"]}  |  {r["company"]}', bold=True)
        p.paragraph_format.space_before = Pt(6)
        para(r["dates"], italic=True)
        if r["intro"]:
            para(humanize(r["intro"]))
        for b in r["bullets"]:
            para(humanize(b), style="List Bullet")
    heading("Technical Skills")
    for g in cv["skills"]:
        p = doc.add_paragraph()
        p.add_run(f'{g["name"]}: ').bold = True
        p.add_run(humanize(", ".join(g["items"])))
    heading("Honours & Recognition")
    for x in cv["recognition"]:
        para(humanize(x), style="List Bullet")
    if education:
        heading("Education")
        for e in education:
            para(f'{e.get("qualification")}, {e.get("institution")}')
    if links:
        heading("Links")
        for l in links:
            para(humanize(l))
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))


def build_cv_variants(force: bool = False) -> Dict[str, Dict]:
    """Write one tailored CV per job type into data/cv_variants/<lane>/. Cached: a variant is rebuilt only when its
    sources, focus or CV_VARIANT_VERSION change (and the model call itself is answered from the AI cache)."""
    from src.services.ai_service import AIService
    from src.services.form_engine.answerer import WRITING_RULES
    from src.services.cv_data import APPLICATION_PROFILE
    source = _cv_sources()
    base = _read_cv(settings.BASE_DIR / settings.CV_FILES["applied_ai"]).splitlines()
    # name and contact line as on the current CV, with the application email everywhere
    header = [base[0].strip(), re.sub(r"[\w.+-]+@[\w-]+\.[\w.]+", APPLICATION_PROFILE["email"], base[2].strip())]
    links = []
    if "LINKS" in base:
        links = [l for l in base[base.index("LINKS") + 1:] if l.strip()]
    system = f"""You rewrite a senior engineer's CV for one type of job. The SOURCES below are the only facts you may use.

RULES
- Truth only: every employer, title, date, number, award and technology must come from the SOURCES. Never add a metric,
  employer, title, date or skill that is not there. You may reorder, select, merge and reword.
- Experience: start with deepPBR.io (CANDIDATE DATA current_work: role "Founder, Architect & Sole Engineer", "June 2026 - Present"),
  lead its bullets with the models he designed and trained himself (own_models), then list the roles from CV A in order with
  CV A's exact titles, companies and dates (CV B may supply extra true bullets for the same roles).
- Tailor to the JOB TYPE in the prompt: headline, summary, which specialisations, bullet order and skill groups favour it.
  Keep everything relevant; drop only what clearly does not serve this job type.
- Length: about two pages (750-950 words). 4-6 bullets for deepPBR.io and PLAiR, 2-4 for older roles.
- Skills: 4-6 groups, plain comma lists, strongest for this job type first.
- Recognition: awards, grants, programmes and talks from the SOURCES.

{WRITING_RULES}

SOURCES
{source}"""
    ai, results = AIService(), {}
    settings.CV_VARIANT_DIR.mkdir(parents=True, exist_ok=True)
    for lane, focus in settings.CV_VARIANTS.items():
        out = settings.CV_VARIANT_DIR / lane / CV_FILE_NAME
        digest = hashlib.sha256(f"{CV_VARIANT_VERSION}|{focus}|{system}".encode("utf-8")).hexdigest()[:16]
        meta_path = out.with_suffix(".json")
        authored = out.parent / "cv.json"       # a hand-written CV for this job type wins over generation
        if authored.exists():
            cv = json.loads(authored.read_text(encoding="utf-8"))
            issues = _unsupported(cv, source)
            _render_cv(cv, header, APPLICATION_PROFILE.get("education", []), links, out)
            results[lane] = {"digest": "authored", "file": str(out), "issues": issues, "headline": cv["headline"], "words": len(_read_cv(out).split())}
            meta_path.write_text(json.dumps(results[lane], indent=1, ensure_ascii=False), encoding="utf-8")
            custom_print("SUCCESS" if not issues else "WARNING", f"CV ({lane}, written): {out}" + (f" - check: {issues[:3]}" if issues else ""))
            continue
        if not force and out.exists() and meta_path.exists() and json.loads(meta_path.read_text(encoding="utf-8")).get("digest") == digest:
            results[lane] = json.loads(meta_path.read_text(encoding="utf-8"))
            custom_print("INFO", f"CV ({lane}): unchanged, kept {out}")
            continue
        prompt = f"JOB TYPE: {focus}\n\nWrite the CV for this job type."
        cv = asyncio.run(ai.structured(prompt, CVVariant, role="pro", temperature=0.4, system=system))
        issues = _unsupported(cv, source) if cv else ["no answer from the model"]
        if cv and issues:
            fix = prompt + "\n\nYour previous draft used facts that are not in the SOURCES - remove or correct them:\n- " + "\n- ".join(issues[:20])
            cv = asyncio.run(ai.structured(fix, CVVariant, role="pro", temperature=0.2, system=system))
            issues = _unsupported(cv, source) if cv else ["no answer from the model"]
        if not cv:
            custom_print("ERROR", f"CV ({lane}): generation failed")
            continue
        _render_cv(cv, header, APPLICATION_PROFILE.get("education", []), links, out)
        results[lane] = {"digest": digest, "file": str(out), "issues": issues, "headline": cv["headline"], "words": len(_read_cv(out).split())}
        meta_path.write_text(json.dumps(results[lane], indent=1, ensure_ascii=False), encoding="utf-8")
        custom_print("SUCCESS" if not issues else "WARNING", f"CV ({lane}): {out}" + (f" - check: {issues[:3]}" if issues else ""))
    return results
